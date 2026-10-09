"""
Métricas del núcleo de IA (sección 9.1 del reto), con numerador, denominador y fallos.

    cd backend && python -m eval.metricas                         # planilla por defecto
    python -m eval.metricas copia_ana.xlsx copia_luis.xlsx         # dos personas: + acuerdo

Sin etiquetas calcula lo automático: cobertura de citas de las fichas, eficiencia
(tiempo, tokens, fallos por proveedor) y costo. Con la planilla etiquetada agrega:
macro-F1 de temas (LLM vs. embeddings vs. palabras clave), exactitud de la relación con
Panamá, Precision@5 y @10 del ranking, precisión y recall de la agrupación y validez de
sustento. Lo que no tiene etiquetas se reporta como pendiente, nunca se inventa.

Salidas: eval/resultados/metricas.md (para Notion) y metricas.json.
"""

import json
import statistics
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from sklearn.metrics import cohen_kappa_score, f1_score

from nucleo import artefactos, config, seguridad
from nucleo.baseline import clasificar_tema_baseline
from nucleo.fichas import paquete_evidencia

CARPETA = Path(__file__).parent
BACKEND = CARPETA.parent
ARTEFACTOS = BACKEND / config.CARPETA_ARTEFACTOS
REGISTRO_LLM = BACKEND / config.CARPETA_CACHE / "llm_log.jsonl"
PLANILLA = CARPETA / "planilla_revisada.xlsx"   # la plantilla vacía es planilla_etiquetas.xlsx
RESULTADOS = CARPETA / "resultados"


def _razon(numerador, denominador):
    return {"numerador": numerador, "denominador": denominador,
            "valor": round(numerador / denominador, 3) if denominador else None}


def _texto_razon(r):
    if not r["denominador"]:
        return "pendiente (sin datos)"
    return f"{r['numerador']} de {r['denominador']} ({r['valor']:.0%})"


def _percentil(valores, p):
    orden = sorted(valores)
    return orden[min(len(orden) - 1, int(len(orden) * p))] if orden else None


# ---------------------------------------------------------------------------
# Lo automático
# ---------------------------------------------------------------------------
def metricas_fichas(eventos, fichas):
    por_id = {e["id_evento"]: e for e in eventos}
    emitidas = validas = 0
    oraciones = oraciones_ok = 0
    fallos = []
    for ficha in fichas:
        paquete, _ = paquete_evidencia(por_id[ficha["id_evento"]])
        buenas, malas = seguridad.validar_afirmaciones(ficha.get("afirmaciones") or [], paquete)
        emitidas += len(ficha.get("afirmaciones") or [])
        validas += len(buenas)
        fallos += [f"{ficha['id_caso']}: {m['motivo_descarte']}" for m in malas]
        for parte in ("brief", "guion_45_60s", "copy_digital"):
            for oracion in seguridad.oraciones((ficha.get("borrador") or {}).get(parte) or ""):
                if seguridad.citas_en_texto(oracion):
                    oraciones += 1
                    problemas = seguridad.cifras_sin_respaldo_por_oracion(oracion, paquete)
                    oraciones_ok += not problemas
                    fallos += [f"{ficha['id_caso']} {parte}: {p}" for p in problemas]
    return {
        "fichas": len(fichas),
        "por_metodo": dict(Counter(f["generado"]["metodo"] for f in fichas)),
        "por_modelo": dict(Counter(f["generado"].get("modelo") or "ninguno" for f in fichas)),
        "bloqueadas": sum(1 for f in fichas if f["validacion"]["bloqueada"]),
        "corregidas": sum(1 for f in fichas if f["generado"].get("reintento_por_bloqueo")),
        "cobertura_citas_afirmaciones": _razon(validas, emitidas),
        "oraciones_citadas_con_cifras_respaldadas": _razon(oraciones_ok, oraciones),
        "fallos": fallos,
    }


def metricas_eficiencia():
    if not REGISTRO_LLM.exists():
        return {"pendiente": "no hay registro de llamadas al LLM"}
    filas = [json.loads(l) for l in REGISTRO_LLM.read_text(encoding="utf-8").splitlines() if l.strip()]
    reales = [f for f in filas if not f.get("desde_cache")]
    exitosas = [f for f in reales if "fallo" not in f]
    por_tarea = {}
    for tarea in sorted({f["tarea"] for f in exitosas}):
        de_tarea = [f for f in exitosas if f["tarea"] == tarea]
        tiempos = [f["segundos"] for f in de_tarea]
        tokens = [f.get("tokens_total") or 0 for f in de_tarea]
        por_tarea[tarea] = {"llamadas": len(de_tarea), "segundos_mediana": round(statistics.median(tiempos), 2),
                            "segundos_p95": _percentil(tiempos, 0.95),
                            "tokens_promedio": round(statistics.mean(tokens)) if tokens else 0}
    return {
        "llamadas_reales": len(reales),
        "exitosas": len(exitosas),
        "desde_cache": len(filas) - len(reales),
        "fallos_por_codigo": dict(Counter(f["fallo"] for f in reales if "fallo" in f)),
        "exitosas_por_modelo": dict(Counter(f["modelo"] for f in exitosas)),
        "tokens_totales": sum(f.get("tokens_total") or 0 for f in exitosas),
        "costo_usd": 0,
        "nota_costo": "Planes gratuitos de Gemini y Groq: costo 0. Se reportan los tokens para estimarlo con precios pagos.",
        "por_tarea": por_tarea,
    }


# ---------------------------------------------------------------------------
# Lo que necesita etiquetas
# ---------------------------------------------------------------------------
def leer_hoja(ruta, nombre, clave, etiquetas):
    """{clave: {etiqueta: valor}} solo con las filas que tienen alguna etiqueta."""
    hoja = load_workbook(ruta, read_only=True)[nombre]
    filas = list(hoja.iter_rows(values_only=True))
    columnas = list(filas[0])
    datos = {}
    for fila in filas[1:]:
        registro = dict(zip(columnas, fila))
        valores = {e: (str(registro.get(e)).strip().lower() if registro.get(e) else None) for e in etiquetas}
        if any(valores.values()):
            datos[registro[clave] if isinstance(clave, str) else tuple(registro[c] for c in clave)] = \
                {**valores, "_fila": registro}
    return datos


def metricas_temas(etiquetas, eventos):
    from nucleo.organizar import clasificar_temas_embeddings, embeber
    con_tema = {i: d for i, d in etiquetas.items() if d["tema"]}
    if not con_tema:
        return {"pendiente": "hoja Etiquetas sin la columna tema completa"}
    guardados = artefactos.cargar("temas.json")
    titulos = {n["id_noticia"]: n["titulo"] for e in eventos for n in e["noticias"]}
    ids = [i for i in con_tema if i in titulos]
    verdad = [con_tema[i]["tema"] for i in ids]

    llm = [(guardados.get(i) or {}).get("tema") if (guardados.get(i) or {}).get("version") ==
           config.VERSION_CLASIFICACION else None for i in ids]
    respaldo = [t for t, _ in clasificar_temas_embeddings(embeber(titulos[i] for i in ids))]
    palabras = [clasificar_tema_baseline(titulos[i]) for i in ids]
    temas = list(config.DEFINICION_TEMAS)

    def resultado(pred, nombre):
        pares = [(v, p) for v, p in zip(verdad, pred) if p is not None]
        if not pares:
            return {"pendiente": f"sin predicciones de {nombre}"}
        v, p = zip(*pares)
        return {"n": len(pares), "macro_f1": round(f1_score(v, p, labels=temas, average="macro", zero_division=0), 3),
                "exactitud": _razon(sum(a == b for a, b in pares), len(pares)),
                "fallos": [f"{i}: humano={a} · sistema={b}" for i, a, b in zip(ids, verdad, pred) if b and a != b][:15]}

    relacion = [(con_tema[i]["relacion_panama"], (guardados.get(i) or {}).get("relacion_panama")) for i in ids]
    relacion = [(a, b) for a, b in relacion if a and b]
    return {
        "etiquetados": len(ids),
        "llm": resultado(llm, "LLM"),
        "embeddings": resultado(respaldo, "embeddings"),
        "palabras_clave": resultado(palabras, "palabras clave"),
        "relacion_panama_llm": _razon(sum(a == b for a, b in relacion), len(relacion)),
    }


def metricas_agenda(etiquetas, eventos):
    posicion = {e["id_evento"]: e["prioridad"]["posicion"] for e in eventos}
    con = [(posicion[i], d["en_agenda_tvn"]) for i, d in etiquetas.items() if d["en_agenda_tvn"] and i in posicion]
    if not con:
        return {"pendiente": "hoja Agenda sin etiquetar"}
    def precision(k):
        top = [v for p, v in con if p <= k]
        return _razon(sum(v == "si" for v in top), len(top))
    return {"precision_5": precision(5), "precision_10": precision(10), "precision_15": precision(15),
            "nota": "Exploratoria si quien etiqueta no es editor/a de TVN (sección 9.1)."}


def metricas_pares(etiquetas, eventos):
    evento_de = {n["id_noticia"]: e["id_evento"] for e in eventos for n in e["noticias"]}
    con = [(evento_de.get(a) == evento_de.get(b), d["mismo_evento"] == "si")
           for (a, b), d in etiquetas.items() if d["mismo_evento"]]
    if not con:
        return {"pendiente": "hoja Pares sin etiquetar"}
    juntos = [h for s, h in con if s]
    mismos = [s for s, h in con if h]
    return {"precision": _razon(sum(juntos), len(juntos)), "recall": _razon(sum(mismos), len(mismos)),
            "nota": "Recall sobre la muestra de pares: incluye solo pares separados parecidos (similitud >= 0,60)."}


def metricas_sustento(etiquetas):
    valores = [d["respaldada"] for d in etiquetas.values() if d["respaldada"]]
    if not valores:
        return {"pendiente": "hoja Afirmaciones sin etiquetar"}
    conteo = Counter(valores)
    return {"validez": _razon(conteo["si"], len(valores)), "parcial": conteo["parcial"], "no": conteo["no"],
            "meta": "90 % sobre al menos 30 afirmaciones (sección 9.1)",
            "fallos": [f"{d['_fila']['id_caso']}: {d['_fila']['texto'][:90]}" for d in etiquetas.values()
                       if d["respaldada"] in ("no", "parcial")]}


def acuerdo(planillas):
    """Kappa de Cohen entre las dos primeras planillas, por columna etiquetada."""
    if len(planillas) < 2:
        return {"pendiente": "se necesita una segunda persona (pasar dos planillas)"}
    resultado = {}
    for hoja, clave, columna in [("Etiquetas", "id_noticia", "tema"), ("Etiquetas", "id_noticia", "relacion_panama"),
                                 ("Agenda", "id_evento", "en_agenda_tvn"), ("Pares", ("id_a", "id_b"), "mismo_evento"),
                                 ("Afirmaciones", "n", "respaldada")]:
        a, b = (leer_hoja(p, hoja, clave, [columna]) for p in planillas[:2])
        comunes = [k for k in a if k in b and a[k][columna] and b[k][columna]]
        if comunes:
            resultado[f"{hoja}.{columna}"] = {"n": len(comunes), "kappa": round(cohen_kappa_score(
                [a[k][columna] for k in comunes], [b[k][columna] for k in comunes]), 3)}
    return resultado


def quien_etiqueto(planilla):
    """Quién etiquetó y cómo: lo que dice la primera celda de la planilla después de "—"."""
    titulo = str(load_workbook(planilla)["Instrucciones"]["A1"].value or "")
    return titulo.partition("—")[2].strip() or None


# ---------------------------------------------------------------------------
def calcular(planillas):
    eventos = json.loads((ARTEFACTOS / "eventos.json").read_text(encoding="utf-8"))
    fichas = [json.loads(l) for l in (ARTEFACTOS / "fichas.jsonl").read_text(encoding="utf-8").splitlines()]
    principal = planillas[0]
    return {
        "planillas": [str(p) for p in planillas],
        "referencia": quien_etiqueto(principal),
        "fichas": metricas_fichas(eventos, fichas),
        "eficiencia": metricas_eficiencia(),
        "temas": metricas_temas(leer_hoja(principal, "Etiquetas", "id_noticia",
                                          ["tema", "relacion_panama", "en_agenda_tvn"]), eventos),
        "agenda": metricas_agenda(leer_hoja(principal, "Agenda", "id_evento", ["en_agenda_tvn"]), eventos),
        "agrupacion": metricas_pares(leer_hoja(principal, "Pares", ("id_a", "id_b"), ["mismo_evento"]), eventos),
        "sustento": metricas_sustento(leer_hoja(principal, "Afirmaciones", "n", ["respaldada"])),
        "acuerdo_entre_personas": acuerdo(planillas),
    }


def a_markdown(m):
    f, e, t = m["fichas"], m["eficiencia"], m["temas"]
    l = ["# Métricas del núcleo de IA", "",
         "Calculadas con `python -m eval.metricas` sobre la corrida del paquete real. Numerador, denominador "
         "y fallos visibles: nada se esconde en un promedio.", "",
         "## Fichas y citas", "", "| Métrica | Resultado |", "|---|---|",
         f"| Fichas | {f['fichas']} (por método: {f['por_metodo']}) |",
         f"| Fichas bloqueadas por el validador | {f['bloqueadas']} (rescatadas con corrección: {f['corregidas']}) |",
         f"| Cobertura de citas: afirmaciones con cita válida | {_texto_razon(f['cobertura_citas_afirmaciones'])} |",
         f"| Oraciones citadas del borrador con cifras respaldadas | {_texto_razon(f['oraciones_citadas_con_cifras_respaldadas'])} |",
         ""]
    if "pendiente" not in e:
        l += ["## Eficiencia y costo", "",
              f"{e['llamadas_reales']} llamadas reales al LLM ({e['exitosas']} exitosas) y {e['desde_cache']} "
              f"respondidas desde la caché. Fallos por código: {e['fallos_por_codigo'] or 'ninguno'}. "
              f"Tokens: {e['tokens_totales']:,}. Costo: {e['costo_usd']} USD ({e['nota_costo']})", "",
              "| Tarea | Llamadas | Mediana (s) | p95 (s) | Tokens promedio |", "|---|---|---|---|---|"]
        l += [f"| {k} | {v['llamadas']} | {v['segundos_mediana']} | {v['segundos_p95']} | {v['tokens_promedio']} |"
              for k, v in e["por_tarea"].items()]
        l.append("")

    if m["referencia"]:
        l += [f"> **Referencia:** {m['referencia']}.", ""]
    l += ["## Clasificación de temas (contra las etiquetas de referencia)", ""]
    if "pendiente" in t:
        l.append(f"Pendiente: {t['pendiente']}.")
    else:
        l += [f"{t['etiquetados']} titulares etiquetados.", "", "| Método | Macro-F1 | Exactitud |", "|---|---|---|"]
        for nombre, clave in [("LLM", "llm"), ("Embeddings (respaldo sin LLM)", "embeddings"),
                              ("Palabras clave (baseline)", "palabras_clave")]:
            r = t[clave]
            l.append(f"| {nombre} | {r.get('macro_f1', '—')} | {_texto_razon(r['exactitud']) if 'exactitud' in r else r.get('pendiente')} |")
        l += ["", f"Relación con Panamá (LLM): {_texto_razon(t['relacion_panama_llm'])}."]
    l.append("")

    for titulo, clave, campos in [("Ranking (Precision@k)", "agenda", ["precision_5", "precision_10", "precision_15"]),
                                  ("Agrupación de noticias", "agrupacion", ["precision", "recall"]),
                                  ("Validez de sustento", "sustento", ["validez"])]:
        datos = m[clave]
        l += [f"## {titulo}", ""]
        if "pendiente" in datos:
            l.append(f"Pendiente: {datos['pendiente']}.")
        else:
            l += [f"- {c}: {_texto_razon(datos[c])}" for c in campos]
            if datos.get("nota"):
                l.append(f"- Nota: {datos['nota']}")
        l.append("")

    a = m["acuerdo_entre_personas"]
    l += ["## Acuerdo entre personas (kappa de Cohen)", ""]
    l += [f"Pendiente: {a['pendiente']}."] if "pendiente" in a else [f"- {k}: {v['kappa']} (n={v['n']})" for k, v in a.items()]
    return "\n".join(l) + "\n"


def main():
    planillas = [Path(p) for p in sys.argv[1:]] or [PLANILLA]
    m = calcular(planillas)
    RESULTADOS.mkdir(exist_ok=True)
    (RESULTADOS / "metricas.json").write_text(json.dumps(m, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    texto = a_markdown(m)
    (RESULTADOS / "metricas.md").write_text(texto, encoding="utf-8")
    print(texto)


if __name__ == "__main__":
    main()
