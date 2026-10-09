"""
Crea la planilla para etiquetar a mano una muestra de titulares reales.

    cd backend && python -m eval.crear_planilla

Las etiquetas humanas son la referencia para medir el sistema (sección 9.1 del
reto): macro-F1 de temas (LLM vs. embeddings vs. palabras clave), exactitud de la
relación con Panamá y Precision@5 del ranking.

Reglas para que la medición valga:
- La planilla NO muestra lo que respondió el sistema (etiquetado a ciegas).
- La muestra es al azar con semilla fija: se puede regenerar igual.
- Idealmente dos personas etiquetan por separado (cada una en su copia); así se
  mide también cuánto coinciden entre ellas.
"""

import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from nucleo import config
from nucleo.organizar import filtrar_por_idioma
from nucleo.pipeline import cargar_desde_paquete

CARPETA = Path(__file__).parent
DATOS = CARPETA.parent.parent / "datos"
SALIDA = CARPETA / "planilla_etiquetas.xlsx"
TAMANO_MUESTRA = 100
SEMILLA = 42

COLUMNAS = [
    # (título, ancho, la completa la persona)
    ("n", 5, False),
    ("id_noticia", 17, False),
    ("titulo", 70, False),
    ("medio", 22, False),
    ("idioma", 7, False),
    ("tema", 20, True),
    ("relacion_panama", 17, True),
    ("en_agenda_tvn", 15, True),
    ("notas", 40, True),
]


def _instrucciones(hoja):
    texto = [
        ("Planilla de etiquetado · núcleo de IA", True),
        ("", False),
        ("Para qué sirve: medir con datos reales si el sistema clasifica bien. Sus respuestas son la "
         "referencia contra la que se compara el LLM.", False),
        ("Cómo completarla:", True),
        ("1. Lea solo el titular (no busque la nota): el sistema tampoco ve más que eso.", False),
        ("2. En la hoja Etiquetas complete las columnas amarillas usando los menús desplegables.", False),
        ("3. No consulte las respuestas del sistema antes de terminar.", False),
        ("4. Si duda entre un tema y \"otro\", elija \"otro\" y explique en notas.", False),
        ("5. Si etiquetan dos personas, cada una usa su propia copia del archivo.", False),
        ("", False),
        ("Columna tema", True),
        *[(f"  {t}: {d}", False) for t, d in config.DEFINICION_TEMAS.items()],
        ("", False),
        ("Columna relacion_panama", True),
        *[(f"  {r}: {d}", False) for r, d in config.RELACIONES_PANAMA.items()],
        ("", False),
        ("Columna en_agenda_tvn", True),
        ("  si: un editor de TVN debería considerarla para la agenda del día.", False),
        ("  no: no merece revisión editorial hoy (irrelevante, de nicho, entretenimiento...).", False),
    ]
    for fila, (linea, negrita) in enumerate(texto, start=1):
        celda = hoja.cell(row=fila, column=1, value=linea)
        celda.font = Font(bold=negrita, size=12 if fila == 1 else 11)
        celda.alignment = Alignment(wrap_text=True, vertical="top")
    hoja.column_dimensions["A"].width = 120


def _etiquetas(hoja, muestra):
    amarillo = PatternFill("solid", fgColor="FFF2CC")
    for col, (titulo, ancho, a_completar) in enumerate(COLUMNAS, start=1):
        celda = hoja.cell(row=1, column=col, value=titulo)
        celda.font = Font(bold=True)
        if a_completar:
            celda.fill = amarillo
        hoja.column_dimensions[celda.column_letter].width = ancho
    hoja.freeze_panes = "A2"

    for fila, n in enumerate(muestra, start=2):
        valores = [fila - 1, n["id_noticia"], n["titulo"], n["medio"], n.get("idioma") or ""]
        for col, valor in enumerate(valores, start=1):
            hoja.cell(row=fila, column=col, value=valor).alignment = Alignment(wrap_text=True, vertical="top")
        for col in range(len(valores) + 1, len(COLUMNAS) + 1):
            hoja.cell(row=fila, column=col).fill = amarillo

    ultima = len(muestra) + 1
    menus = {
        "F": list(config.DEFINICION_TEMAS),
        "G": list(config.RELACIONES_PANAMA),
        "H": ["si", "no"],
    }
    for columna, opciones in menus.items():
        menu = DataValidation(type="list", formula1='"' + ",".join(opciones) + '"', allow_blank=True,
                              showErrorMessage=True, errorTitle="Valor no válido",
                              error="Elija una opción del menú desplegable.")
        hoja.add_data_validation(menu)
        menu.add(f"{columna}2:{columna}{ultima}")


def crear(tamano=TAMANO_MUESTRA, semilla=SEMILLA, salida=SALIDA):
    noticias, *_ = cargar_desde_paquete(DATOS)
    en_panel, _ = filtrar_por_idioma(noticias)
    en_panel.sort(key=lambda n: n["id_noticia"])  # orden fijo antes de sortear
    muestra = random.Random(semilla).sample(en_panel, min(tamano, len(en_panel)))

    libro = Workbook()
    _instrucciones(libro.active)
    libro.active.title = "Instrucciones"
    _etiquetas(libro.create_sheet("Etiquetas"), muestra)
    libro.save(salida)
    return salida, len(muestra), len(en_panel)


if __name__ == "__main__":
    ruta, n, total = crear()
    print(f"Planilla creada: {ruta} ({n} titulares sorteados de {total} en el panel, semilla {SEMILLA})")
