"""
Pruebas de eval/metricas.py con una planilla etiquetada de forma simulada.
"""

import json
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from eval import metricas
from nucleo import artefactos

PLANILLA = Path(__file__).parent.parent / "eval" / "planilla_etiquetas.xlsx"
EVENTOS = json.loads((metricas.ARTEFACTOS / "eventos.json").read_text(encoding="utf-8"))

pytestmark = pytest.mark.skipif(not PLANILLA.exists(), reason="falta la planilla")


def _completar(ruta, hoja, columna, valor_de):
    libro = load_workbook(ruta)
    h = libro[hoja]
    columnas = [c.value for c in h[1]]
    for fila in h.iter_rows(min_row=2):
        registro = dict(zip(columnas, [c.value for c in fila]))
        fila[columnas.index(columna)].value = valor_de(registro)
    libro.save(ruta)


def test_sin_etiquetas_todo_queda_pendiente_y_lo_automatico_se_calcula():
    m = metricas.calcular([PLANILLA])
    assert m["fichas"]["cobertura_citas_afirmaciones"]["denominador"] > 0
    for clave in ("temas", "agenda", "agrupacion", "sustento", "acuerdo_entre_personas"):
        assert "pendiente" in m[clave], clave


def test_con_etiquetas_simuladas(tmp_path):
    copia = tmp_path / "etiquetada.xlsx"
    shutil.copy(PLANILLA, copia)
    temas = artefactos.cargar("temas.json")
    evento_de = {n["id_noticia"]: e["id_evento"] for e in EVENTOS for n in e["noticias"]}

    # Temas iguales a los del LLM -> el LLM debe sacar macro-F1 1
    _completar(copia, "Etiquetas", "tema", lambda r: (temas.get(r["id_noticia"]) or {}).get("tema", "otro"))
    # Todo en la agenda -> Precision@5 = 1
    _completar(copia, "Agenda", "en_agenda_tvn", lambda r: "si")
    # La persona coincide con el sistema en los pares -> precisión y recall = 1
    _completar(copia, "Pares", "mismo_evento",
               lambda r: "si" if evento_de.get(r["id_a"]) == evento_de.get(r["id_b"]) else "no")
    # Una afirmación sin respaldo -> validez < 1 y aparece en los fallos
    _completar(copia, "Afirmaciones", "respaldada", lambda r: "no" if r["n"] == 1 else "si")

    m = metricas.calcular([copia, copia])
    assert m["temas"]["llm"]["macro_f1"] == 1.0
    assert m["temas"]["palabras_clave"]["macro_f1"] <= 1.0
    assert m["agenda"]["precision_5"]["valor"] == 1.0
    assert m["agrupacion"]["precision"]["valor"] == 1.0 and m["agrupacion"]["recall"]["valor"] == 1.0
    assert m["sustento"]["validez"]["valor"] < 1.0 and len(m["sustento"]["fallos"]) == 1
    assert m["acuerdo_entre_personas"]["Etiquetas.tema"]["kappa"] == 1.0
    assert "Macro-F1" in metricas.a_markdown(m)
